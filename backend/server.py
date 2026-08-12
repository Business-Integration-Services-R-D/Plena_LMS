from fastapi import FastAPI, APIRouter, HTTPException, Depends, Request, Response, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import re
import logging
import asyncio
import httpx
import aiofiles
from pathlib import Path
from pydantic import BaseModel, Field
from typing import List, Optional, Any
import uuid
from datetime import datetime, timezone, timedelta
from io import BytesIO
import unicodedata

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

VIDEO_DIR = ROOT_DIR / 'uploads' / 'videos'
VIDEO_DIR.mkdir(parents=True, exist_ok=True)
MAX_VIDEO_BYTES = 250 * 1024 * 1024
ADMIN_SEED_EMAIL = "tugberkkalay@gmail.com"
AUTH_BYPASS = os.environ.get("AUTH_BYPASS", "").lower() in ("1", "true", "yes")
IS_PROD = os.environ.get("ENV", "development") == "production"

app = FastAPI()
api_router = APIRouter(prefix="/api")
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


def set_session_cookie(response: Response, session_token: str):
    response.set_cookie(
        "session_token",
        session_token,
        max_age=7 * 24 * 3600,
        httponly=True,
        secure=IS_PROD,
        samesite="none" if IS_PROD else "lax",
        path="/",
    )


async def create_user_session(user: dict, response: Response, session_token: Optional[str] = None):
    token = session_token or f"dev_{uuid.uuid4().hex}"
    await db.user_sessions.insert_one({
        "session_id": new_id("sess"),
        "user_id": user["user_id"],
        "session_token": token,
        "expires_at": (datetime.now(timezone.utc) + timedelta(days=7)).isoformat(),
        "created_at": now_iso(),
    })
    set_session_cookie(response, token)
    return token


def now_iso():
    return datetime.now(timezone.utc).isoformat()


def new_id(prefix):
    return f"{prefix}_{uuid.uuid4().hex[:12]}"


def parse_dt(v):
    if isinstance(v, str):
        dt = datetime.fromisoformat(v.replace('Z', '+00:00'))
    else:
        dt = v
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt


async def log_email(to, subject, body, etype):
    await db.email_logs.insert_one({
        "email_id": new_id("mail"), "to": to, "subject": subject,
        "body": body, "type": etype, "created_at": now_iso()
    })


async def log_event(assignment_id, user_id, etype, position=None, data=None):
    await db.events.insert_one({
        "event_id": new_id("evt"), "assignment_id": assignment_id, "user_id": user_id,
        "type": etype, "position": position, "data": data, "created_at": now_iso()
    })


# ---------- AUTH ----------
async def get_current_user(request: Request):
    token = request.cookies.get("session_token")
    if not token:
        auth = request.headers.get("Authorization", "")
        if auth.startswith("Bearer "):
            token = auth.split(" ", 1)[1]
    if not token:
        raise HTTPException(status_code=401, detail="Oturum bulunamadı")
    session = await db.user_sessions.find_one({"session_token": token}, {"_id": 0})
    if not session:
        raise HTTPException(status_code=401, detail="Geçersiz oturum")
    expires_at = parse_dt(session["expires_at"])
    if expires_at < datetime.now(timezone.utc):
        raise HTTPException(status_code=401, detail="Oturum süresi doldu")
    user = await db.users.find_one({"user_id": session["user_id"]}, {"_id": 0})
    if not user:
        raise HTTPException(status_code=401, detail="Kullanıcı bulunamadı")
    return user


async def require_admin(user=Depends(get_current_user)):
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Bu işlem için admin yetkisi gerekli")
    return user


class SessionRequest(BaseModel):
    session_id: str


@api_router.post("/auth/session")
async def create_session(body: SessionRequest, response: Response):
    async with httpx.AsyncClient() as hc:
        r = await hc.get(
            "https://demobackend.emergentagent.com/auth/v1/env/oauth/session-data",
            headers={"X-Session-ID": body.session_id}, timeout=15
        )
    if r.status_code != 200:
        raise HTTPException(status_code=401, detail="Google oturumu doğrulanamadı")
    data = r.json()
    email = data["email"].lower().strip()
    user = await db.users.find_one({"email": email}, {"_id": 0})
    if not user:
        raise HTTPException(status_code=403, detail="Bu e-posta için tanımlı bir kullanıcı yok. Lütfen yöneticinizle iletişime geçin.")
    await db.users.update_one({"email": email}, {"$set": {
        "name": data.get("name") or user.get("name"),
        "picture": data.get("picture"), "status": "active", "last_login": now_iso()
    }})
    user = await db.users.find_one({"email": email}, {"_id": 0})
    await create_user_session(user, response, session_token=data["session_token"])
    return user


@api_router.post("/auth/dev-login")
async def dev_login(response: Response, role: str = "admin"):
    """Local-only bypass: create a session without Google OAuth."""
    if not AUTH_BYPASS:
        raise HTTPException(status_code=404, detail="Not found")
    if role not in ("admin", "employee"):
        raise HTTPException(status_code=400, detail="Geçersiz rol")
    email = ADMIN_SEED_EMAIL if role == "admin" else "dev.employee@localhost"
    user = await db.users.find_one({"email": email}, {"_id": 0})
    if not user:
        user_doc = {
            "user_id": new_id("user"),
            "email": email,
            "name": "Local Admin" if role == "admin" else "Local Employee",
            "role": role,
            "status": "active",
            "picture": None,
            "created_at": now_iso(),
        }
        await db.users.insert_one(user_doc)
        user = {k: v for k, v in user_doc.items()}
    else:
        await db.users.update_one(
            {"email": email},
            {"$set": {"role": role, "status": "active", "last_login": now_iso()}},
        )
        user = await db.users.find_one({"email": email}, {"_id": 0})
    await create_user_session(user, response)
    return user


@api_router.get("/auth/me")
async def auth_me(user=Depends(get_current_user)):
    return user


@api_router.post("/auth/logout")
async def logout(request: Request, response: Response, user=Depends(get_current_user)):
    token = request.cookies.get("session_token")
    if token:
        await db.user_sessions.delete_one({"session_token": token})
    response.delete_cookie("session_token", path="/")
    return {"ok": True}


# ---------- USERS & GROUPS ----------
class UserCreate(BaseModel):
    email: str
    name: str
    role: str = "employee"


class UserUpdate(BaseModel):
    name: Optional[str] = None
    role: Optional[str] = None


@api_router.get("/users")
async def list_users(admin=Depends(require_admin)):
    return await db.users.find({}, {"_id": 0}).sort("created_at", -1).to_list(1000)


@api_router.post("/users")
async def create_user(body: UserCreate, admin=Depends(require_admin)):
    email = body.email.lower().strip()
    if not re.match(r"[^@]+@[^@]+\.[^@]+", email):
        raise HTTPException(status_code=400, detail="Geçersiz e-posta adresi")
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Bu e-posta zaten kayıtlı")
    doc = {"user_id": new_id("user"), "email": email, "name": body.name, "role": body.role,
           "status": "invited", "picture": None, "created_at": now_iso()}
    await db.users.insert_one(doc)
    await log_email(email, "Plena LMS'a Davet Edildiniz",
                    f"Merhaba {body.name}, Plena LMS eğitim platformuna davet edildiniz. Google hesabınızla ({email}) giriş yaparak eğitimlerinize erişebilirsiniz.",
                    "activation")
    doc.pop("_id", None)
    return doc


@api_router.put("/users/{user_id}")
async def update_user(user_id: str, body: UserUpdate, admin=Depends(require_admin)):
    updates = {k: v for k, v in body.model_dump().items() if v is not None}
    await db.users.update_one({"user_id": user_id}, {"$set": updates})
    return await db.users.find_one({"user_id": user_id}, {"_id": 0})


@api_router.delete("/users/{user_id}")
async def delete_user(user_id: str, admin=Depends(require_admin)):
    user = await db.users.find_one({"user_id": user_id}, {"_id": 0})
    if user and user["email"] == ADMIN_SEED_EMAIL:
        raise HTTPException(status_code=400, detail="Ana admin silinemez")
    await db.users.delete_one({"user_id": user_id})
    await db.user_sessions.delete_many({"user_id": user_id})
    await db.groups.update_many({}, {"$pull": {"member_ids": user_id}})
    return {"ok": True}


@api_router.post("/users/{user_id}/resend-activation")
async def resend_activation(user_id: str, admin=Depends(require_admin)):
    user = await db.users.find_one({"user_id": user_id}, {"_id": 0})
    if not user:
        raise HTTPException(status_code=404, detail="Kullanıcı bulunamadı")
    await log_email(user["email"], "Plena LMS Aktivasyon Hatırlatması",
                    f"Merhaba {user['name']}, Plena LMS hesabınız sizi bekliyor. Google hesabınızla giriş yapabilirsiniz.",
                    "activation")
    return {"ok": True}


class GroupCreate(BaseModel):
    name: str
    member_ids: List[str] = []


@api_router.get("/groups")
async def list_groups(admin=Depends(require_admin)):
    return await db.groups.find({}, {"_id": 0}).sort("created_at", -1).to_list(500)


@api_router.post("/groups")
async def create_group(body: GroupCreate, admin=Depends(require_admin)):
    doc = {"group_id": new_id("grp"), "name": body.name, "member_ids": body.member_ids, "created_at": now_iso()}
    await db.groups.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api_router.put("/groups/{group_id}")
async def update_group(group_id: str, body: GroupCreate, admin=Depends(require_admin)):
    await db.groups.update_one({"group_id": group_id}, {"$set": {"name": body.name, "member_ids": body.member_ids}})
    return await db.groups.find_one({"group_id": group_id}, {"_id": 0})


@api_router.delete("/groups/{group_id}")
async def delete_group(group_id: str, admin=Depends(require_admin)):
    await db.groups.delete_one({"group_id": group_id})
    return {"ok": True}


# ---------- QUESTIONS ----------
class QuestionBody(BaseModel):
    text: str
    qtype: str  # multiple_choice | free_text
    options: List[str] = []
    correct_index: Optional[int] = None
    category: Optional[str] = None


@api_router.get("/questions")
async def list_questions(admin=Depends(require_admin)):
    return await db.questions.find({}, {"_id": 0}).sort("created_at", -1).to_list(2000)


@api_router.post("/questions")
async def create_question(body: QuestionBody, admin=Depends(require_admin)):
    if body.qtype == "multiple_choice" and (len(body.options) < 2 or body.correct_index is None):
        raise HTTPException(status_code=400, detail="Çoktan seçmeli soruda en az 2 seçenek ve doğru cevap gerekli")
    doc = {"question_id": new_id("q"), **body.model_dump(), "created_at": now_iso()}
    await db.questions.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api_router.put("/questions/{question_id}")
async def update_question(question_id: str, body: QuestionBody, admin=Depends(require_admin)):
    await db.questions.update_one({"question_id": question_id}, {"$set": body.model_dump()})
    return await db.questions.find_one({"question_id": question_id}, {"_id": 0})


@api_router.delete("/questions/{question_id}")
async def delete_question(question_id: str, admin=Depends(require_admin)):
    await db.questions.delete_one({"question_id": question_id})
    return {"ok": True}


# ---------- TRAININGS ----------
class TrainingCreate(BaseModel):
    title: str
    description: Optional[str] = ""


class TrainingUpdate(BaseModel):
    title: Optional[str] = None
    description: Optional[str] = None
    checkpoints: Optional[List[dict]] = None
    quiz: Optional[dict] = None
    duration: Optional[float] = None


@api_router.get("/trainings")
async def list_trainings(admin=Depends(require_admin)):
    trainings = await db.trainings.find({}, {"_id": 0}).sort("created_at", -1).to_list(500)
    for t in trainings:
        t["assignment_count"] = await db.assignments.count_documents({"training_id": t["training_id"]})
    return trainings


@api_router.post("/trainings")
async def create_training(body: TrainingCreate, admin=Depends(require_admin)):
    doc = {"training_id": new_id("trn"), "title": body.title, "description": body.description,
           "video_filename": None, "video_size": None, "duration": None,
           "checkpoints": [], "quiz": None, "created_at": now_iso()}
    await db.trainings.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api_router.get("/trainings/{training_id}")
async def get_training(training_id: str, admin=Depends(require_admin)):
    t = await db.trainings.find_one({"training_id": training_id}, {"_id": 0})
    if not t:
        raise HTTPException(status_code=404, detail="Eğitim bulunamadı")
    return t


@api_router.put("/trainings/{training_id}")
async def update_training(training_id: str, body: TrainingUpdate, admin=Depends(require_admin)):
    updates = {k: v for k, v in body.model_dump().items() if v is not None}
    if "checkpoints" in updates:
        for cp in updates["checkpoints"]:
            if not cp.get("id"):
                cp["id"] = new_id("cp")
        updates["checkpoints"] = sorted(updates["checkpoints"], key=lambda c: c["time"])
    await db.trainings.update_one({"training_id": training_id}, {"$set": updates})
    return await db.trainings.find_one({"training_id": training_id}, {"_id": 0})


@api_router.delete("/trainings/{training_id}")
async def delete_training(training_id: str, admin=Depends(require_admin)):
    t = await db.trainings.find_one({"training_id": training_id}, {"_id": 0})
    if t and t.get("video_filename"):
        try:
            (VIDEO_DIR / t["video_filename"]).unlink(missing_ok=True)
        except Exception:
            pass
    await db.trainings.delete_one({"training_id": training_id})
    await db.assignments.delete_many({"training_id": training_id})
    return {"ok": True}


@api_router.post("/trainings/{training_id}/video")
async def upload_video(training_id: str, file: UploadFile = File(...), duration: float = Form(0), admin=Depends(require_admin)):
    t = await db.trainings.find_one({"training_id": training_id}, {"_id": 0})
    if not t:
        raise HTTPException(status_code=404, detail="Eğitim bulunamadı")
    if not (file.filename or "").lower().endswith(".mp4"):
        raise HTTPException(status_code=400, detail="Sadece MP4 formatı destekleniyor")
    filename = f"{training_id}.mp4"
    path = VIDEO_DIR / filename
    size = 0
    async with aiofiles.open(path, "wb") as f:
        while chunk := await file.read(1024 * 1024):
            size += len(chunk)
            if size > MAX_VIDEO_BYTES:
                await f.close()
                path.unlink(missing_ok=True)
                raise HTTPException(status_code=400, detail="Video 250MB sınırını aşıyor")
            await f.write(chunk)
    await db.trainings.update_one({"training_id": training_id}, {"$set": {
        "video_filename": filename, "video_size": size, "duration": duration or t.get("duration")
    }})
    return {"ok": True, "size": size}


@api_router.get("/videos/{training_id}")
async def stream_video(training_id: str, request: Request, user=Depends(get_current_user)):
    t = await db.trainings.find_one({"training_id": training_id}, {"_id": 0})
    if not t or not t.get("video_filename"):
        raise HTTPException(status_code=404, detail="Video bulunamadı")
    path = VIDEO_DIR / t["video_filename"]
    if not path.exists():
        raise HTTPException(status_code=404, detail="Video dosyası bulunamadı")
    file_size = path.stat().st_size
    range_header = request.headers.get("range")
    start, end = 0, file_size - 1
    status_code = 200
    if range_header:
        m = re.match(r"bytes=(\d+)-(\d*)", range_header)
        if m:
            start = int(m.group(1))
            if m.group(2):
                end = int(m.group(2))
            status_code = 206
    chunk_size = end - start + 1

    def iterfile():
        with open(path, "rb") as f:
            f.seek(start)
            remaining = chunk_size
            while remaining > 0:
                data = f.read(min(1024 * 512, remaining))
                if not data:
                    break
                remaining -= len(data)
                yield data

    headers = {
        "Accept-Ranges": "bytes",
        "Content-Length": str(chunk_size),
        "Content-Type": "video/mp4",
    }
    if status_code == 206:
        headers["Content-Range"] = f"bytes {start}-{end}/{file_size}"
    return StreamingResponse(iterfile(), status_code=status_code, headers=headers)


# ---------- ASSIGNMENTS ----------
class AssignmentCreate(BaseModel):
    training_id: str
    user_ids: List[str] = []
    group_ids: List[str] = []
    start_at: Optional[str] = None
    due_at: Optional[str] = None
    reminder_days: int = 0


@api_router.post("/assignments")
async def create_assignments(body: AssignmentCreate, admin=Depends(require_admin)):
    training = await db.trainings.find_one({"training_id": body.training_id}, {"_id": 0})
    if not training:
        raise HTTPException(status_code=404, detail="Eğitim bulunamadı")
    if not training.get("video_filename"):
        raise HTTPException(status_code=400, detail="Bu eğitime henüz video yüklenmemiş")
    target_ids = set(body.user_ids)
    for gid in body.group_ids:
        g = await db.groups.find_one({"group_id": gid}, {"_id": 0})
        if g:
            target_ids.update(g.get("member_ids", []))
    if not target_ids:
        raise HTTPException(status_code=400, detail="En az bir kullanıcı veya grup seçin")
    start_at = body.start_at or now_iso()
    created = []
    for uid in target_ids:
        user = await db.users.find_one({"user_id": uid}, {"_id": 0})
        if not user:
            continue
        existing = await db.assignments.find_one({"training_id": body.training_id, "user_id": uid, "status": {"$ne": "completed"}})
        if existing:
            continue
        doc = {"assignment_id": new_id("asg"), "training_id": body.training_id, "user_id": uid,
               "status": "assigned", "start_at": start_at, "due_at": body.due_at,
               "reminder_days": body.reminder_days, "last_reminder_at": None,
               "assigned_by": admin["user_id"], "created_at": now_iso(), "completed_at": None}
        await db.assignments.insert_one(doc)
        doc.pop("_id", None)
        created.append(doc)
        await log_email(user["email"], f"Yeni Eğitim Atandı: {training['title']}",
                        f"Merhaba {user['name']}, size '{training['title']}' eğitimi atandı. Başlangıç: {start_at[:10]}." + (f" Son tarih: {body.due_at[:10]}." if body.due_at else ""),
                        "assignment")
    return {"created": len(created), "assignments": created}


@api_router.get("/assignments")
async def list_assignments(admin=Depends(require_admin)):
    assignments = await db.assignments.find({}, {"_id": 0}).sort("created_at", -1).to_list(2000)
    users = {u["user_id"]: u for u in await db.users.find({}, {"_id": 0}).to_list(1000)}
    trainings = {t["training_id"]: t for t in await db.trainings.find({}, {"_id": 0}).to_list(500)}
    for a in assignments:
        u = users.get(a["user_id"], {})
        t = trainings.get(a["training_id"], {})
        a["user_name"] = u.get("name", "-")
        a["user_email"] = u.get("email", "-")
        a["training_title"] = t.get("title", "-")
    return assignments


@api_router.delete("/assignments/{assignment_id}")
async def delete_assignment(assignment_id: str, admin=Depends(require_admin)):
    await db.assignments.delete_one({"assignment_id": assignment_id})
    await db.progress.delete_many({"assignment_id": assignment_id})
    return {"ok": True}


# ---------- EMPLOYEE LEARN ----------
def strip_question(q):
    return {"question_id": q["question_id"], "text": q["text"], "qtype": q["qtype"],
            "options": q.get("options", [])}


async def get_progress(assignment_id, user_id, training_id):
    prog = await db.progress.find_one({"assignment_id": assignment_id}, {"_id": 0})
    if not prog:
        prog = {"progress_id": new_id("prg"), "assignment_id": assignment_id, "user_id": user_id,
                "training_id": training_id, "current_position": 0.0, "max_position": 0.0,
                "watched_seconds": 0.0, "checkpoints_passed": [], "checkpoint_attempts": [],
                "video_completed": False, "quiz_attempts": [], "last_heartbeat": now_iso(),
                "created_at": now_iso()}
        await db.progress.insert_one(dict(prog))
        prog.pop("_id", None)
    return prog


@api_router.get("/my/assignments")
async def my_assignments(user=Depends(get_current_user)):
    now = now_iso()
    assignments = await db.assignments.find(
        {"user_id": user["user_id"], "start_at": {"$lte": now}}, {"_id": 0}
    ).sort("created_at", -1).to_list(200)
    result = []
    for a in assignments:
        t = await db.trainings.find_one({"training_id": a["training_id"]}, {"_id": 0})
        if not t:
            continue
        prog = await db.progress.find_one({"assignment_id": a["assignment_id"]}, {"_id": 0})
        duration = t.get("duration") or 0
        watch_pct = round(min(100, (prog["max_position"] / duration * 100)) if prog and duration else 0)
        result.append({**a, "training_title": t["title"], "training_description": t.get("description", ""),
                       "duration": duration, "has_quiz": bool(t.get("quiz") and t["quiz"].get("question_ids")),
                       "watch_pct": watch_pct,
                       "video_completed": prog.get("video_completed", False) if prog else False})
    return result


@api_router.get("/learn/{assignment_id}")
async def learn_detail(assignment_id: str, user=Depends(get_current_user)):
    a = await db.assignments.find_one({"assignment_id": assignment_id, "user_id": user["user_id"]}, {"_id": 0})
    if not a:
        raise HTTPException(status_code=404, detail="Atama bulunamadı")
    if parse_dt(a["start_at"]) > datetime.now(timezone.utc):
        raise HTTPException(status_code=403, detail="Bu eğitim henüz başlamadı")
    t = await db.trainings.find_one({"training_id": a["training_id"]}, {"_id": 0})
    if not t:
        raise HTTPException(status_code=404, detail="Eğitim bulunamadı")
    prog = await get_progress(assignment_id, user["user_id"], a["training_id"])
    checkpoints = []
    for cp in t.get("checkpoints", []):
        q = await db.questions.find_one({"question_id": cp["question_id"]}, {"_id": 0})
        if q:
            checkpoints.append({"id": cp["id"], "time": cp["time"], "timeout_seconds": cp.get("timeout_seconds", 60),
                                "on_fail": cp.get("on_fail", "start"), "question": strip_question(q)})
    quiz = None
    if t.get("quiz") and t["quiz"].get("question_ids"):
        questions = []
        for qid in t["quiz"]["question_ids"]:
            q = await db.questions.find_one({"question_id": qid}, {"_id": 0})
            if q:
                questions.append(strip_question(q))
        quiz = {"questions": questions, "pass_score": t["quiz"].get("pass_score")}
    return {"assignment": a,
            "training": {"training_id": t["training_id"], "title": t["title"],
                         "description": t.get("description", ""), "duration": t.get("duration"),
                         "checkpoints": checkpoints, "quiz": quiz},
            "progress": {k: prog[k] for k in ["current_position", "max_position", "watched_seconds",
                                              "checkpoints_passed", "video_completed", "quiz_attempts"]}}


class HeartbeatBody(BaseModel):
    position: float
    playing: bool


@api_router.post("/learn/{assignment_id}/heartbeat")
async def heartbeat(assignment_id: str, body: HeartbeatBody, user=Depends(get_current_user)):
    a = await db.assignments.find_one({"assignment_id": assignment_id, "user_id": user["user_id"]}, {"_id": 0})
    if not a:
        raise HTTPException(status_code=404, detail="Atama bulunamadı")
    t = await db.trainings.find_one({"training_id": a["training_id"]}, {"_id": 0})
    prog = await get_progress(assignment_id, user["user_id"], a["training_id"])
    now = datetime.now(timezone.utc)
    elapsed = 3.0
    if prog.get("last_heartbeat"):
        elapsed = min((now - parse_dt(prog["last_heartbeat"])).total_seconds(), 60)
    allowed_max = prog["max_position"] + (elapsed * 1.5 + 3 if body.playing else 3)
    clamped = min(body.position, allowed_max)
    new_max = max(prog["max_position"], clamped)
    # server-side checkpoint gate: cannot progress past an unanswered checkpoint
    for cp in t.get("checkpoints", []):
        if cp["id"] not in prog["checkpoints_passed"] and new_max > cp["time"] + 1.5:
            new_max = cp["time"] + 1.5
            clamped = min(clamped, new_max)
    watched = prog["watched_seconds"] + (min(elapsed, 60) if body.playing else 0)
    updates = {"current_position": clamped, "max_position": new_max,
               "watched_seconds": watched, "last_heartbeat": now.isoformat()}
    if a["status"] == "assigned":
        await db.assignments.update_one({"assignment_id": assignment_id}, {"$set": {"status": "in_progress"}})
    await db.progress.update_one({"assignment_id": assignment_id}, {"$set": updates})
    return {"current_position": clamped, "max_position": new_max}


class EventBody(BaseModel):
    type: str
    position: Optional[float] = None


@api_router.post("/learn/{assignment_id}/event")
async def learn_event(assignment_id: str, body: EventBody, user=Depends(get_current_user)):
    a = await db.assignments.find_one({"assignment_id": assignment_id, "user_id": user["user_id"]}, {"_id": 0})
    if not a:
        raise HTTPException(status_code=404, detail="Atama bulunamadı")
    if body.type not in ["video_started", "video_paused", "video_resumed", "video_closed", "video_ended"]:
        raise HTTPException(status_code=400, detail="Geçersiz olay tipi")
    await log_event(assignment_id, user["user_id"], body.type, body.position)
    return {"ok": True}


class CheckpointAnswer(BaseModel):
    checkpoint_id: str
    answer_index: Optional[int] = None
    answer_text: Optional[str] = None
    timed_out: bool = False


@api_router.post("/learn/{assignment_id}/checkpoint")
async def answer_checkpoint(assignment_id: str, body: CheckpointAnswer, user=Depends(get_current_user)):
    a = await db.assignments.find_one({"assignment_id": assignment_id, "user_id": user["user_id"]}, {"_id": 0})
    if not a:
        raise HTTPException(status_code=404, detail="Atama bulunamadı")
    t = await db.trainings.find_one({"training_id": a["training_id"]}, {"_id": 0})
    cp = next((c for c in t.get("checkpoints", []) if c["id"] == body.checkpoint_id), None)
    if not cp:
        raise HTTPException(status_code=404, detail="Kontrol noktası bulunamadı")
    prog = await get_progress(assignment_id, user["user_id"], a["training_id"])
    q = await db.questions.find_one({"question_id": cp["question_id"]}, {"_id": 0})
    passed = False
    if not body.timed_out:
        if q["qtype"] == "multiple_choice":
            passed = body.answer_index == q.get("correct_index")
        else:
            passed = bool((body.answer_text or "").strip())
    attempt = {"checkpoint_id": cp["id"], "question_id": q["question_id"],
               "answer_index": body.answer_index, "answer_text": body.answer_text,
               "timed_out": body.timed_out, "passed": passed, "at": now_iso()}
    rewind_to = None
    if passed:
        await db.progress.update_one({"assignment_id": assignment_id}, {
            "$addToSet": {"checkpoints_passed": cp["id"]},
            "$push": {"checkpoint_attempts": attempt}})
        await log_event(assignment_id, user["user_id"], "checkpoint_passed", cp["time"], {"checkpoint_id": cp["id"]})
    else:
        if cp.get("on_fail", "start") == "previous":
            prev_times = [c["time"] for c in t.get("checkpoints", [])
                          if c["id"] in prog["checkpoints_passed"] and c["time"] < cp["time"]]
            rewind_to = max(prev_times) if prev_times else 0.0
        else:
            rewind_to = 0.0
        await db.progress.update_one({"assignment_id": assignment_id}, {
            "$set": {"max_position": rewind_to, "current_position": rewind_to},
            "$push": {"checkpoint_attempts": attempt}})
        await log_event(assignment_id, user["user_id"], "checkpoint_failed", cp["time"],
                        {"checkpoint_id": cp["id"], "timed_out": body.timed_out, "rewind_to": rewind_to})
    return {"passed": passed, "rewind_to": rewind_to}


@api_router.post("/learn/{assignment_id}/video-complete")
async def video_complete(assignment_id: str, user=Depends(get_current_user)):
    a = await db.assignments.find_one({"assignment_id": assignment_id, "user_id": user["user_id"]}, {"_id": 0})
    if not a:
        raise HTTPException(status_code=404, detail="Atama bulunamadı")
    t = await db.trainings.find_one({"training_id": a["training_id"]}, {"_id": 0})
    prog = await get_progress(assignment_id, user["user_id"], a["training_id"])
    duration = t.get("duration") or 0
    cp_ids = [c["id"] for c in t.get("checkpoints", [])]
    if duration and prog["max_position"] < duration - 5:
        raise HTTPException(status_code=400, detail="Video henüz tamamen izlenmedi")
    if any(cid not in prog["checkpoints_passed"] for cid in cp_ids):
        raise HTTPException(status_code=400, detail="Tüm kontrol noktaları cevaplanmadı")
    await db.progress.update_one({"assignment_id": assignment_id}, {"$set": {"video_completed": True}})
    has_quiz = bool(t.get("quiz") and t["quiz"].get("question_ids"))
    new_status = "video_completed" if has_quiz else "completed"
    updates = {"status": new_status}
    if not has_quiz:
        updates["completed_at"] = now_iso()
    await db.assignments.update_one({"assignment_id": assignment_id}, {"$set": updates})
    await log_event(assignment_id, user["user_id"], "video_completed", duration)
    return {"ok": True, "has_quiz": has_quiz, "status": new_status}


class QuizSubmit(BaseModel):
    answers: List[dict]  # {question_id, answer_index?, answer_text?}


@api_router.post("/learn/{assignment_id}/quiz")
async def submit_quiz(assignment_id: str, body: QuizSubmit, user=Depends(get_current_user)):
    a = await db.assignments.find_one({"assignment_id": assignment_id, "user_id": user["user_id"]}, {"_id": 0})
    if not a:
        raise HTTPException(status_code=404, detail="Atama bulunamadı")
    t = await db.trainings.find_one({"training_id": a["training_id"]}, {"_id": 0})
    prog = await get_progress(assignment_id, user["user_id"], a["training_id"])
    if not prog.get("video_completed"):
        raise HTTPException(status_code=400, detail="Önce videoyu tamamlamalısınız")
    quiz = t.get("quiz") or {}
    qids = quiz.get("question_ids", [])
    graded = []
    mc_total, mc_correct = 0, 0
    ans_map = {x.get("question_id"): x for x in body.answers}
    for qid in qids:
        q = await db.questions.find_one({"question_id": qid}, {"_id": 0})
        if not q:
            continue
        ans = ans_map.get(qid, {})
        item = {"question_id": qid, "text": q["text"], "qtype": q["qtype"],
                "answer_index": ans.get("answer_index"), "answer_text": ans.get("answer_text")}
        if q["qtype"] == "multiple_choice":
            mc_total += 1
            correct = ans.get("answer_index") == q.get("correct_index")
            if correct:
                mc_correct += 1
            item["correct"] = correct
            item["correct_index"] = q.get("correct_index")
            item["options"] = q.get("options", [])
        else:
            item["correct"] = None
        graded.append(item)
    score = round(mc_correct / mc_total * 100) if mc_total else 100
    pass_score = quiz.get("pass_score")
    passed = score >= pass_score if pass_score else True
    attempt = {"attempt_id": new_id("qa"), "answers": graded, "score": score,
               "passed": passed, "submitted_at": now_iso()}
    await db.progress.update_one({"assignment_id": assignment_id}, {"$push": {"quiz_attempts": attempt}})
    if passed:
        await db.assignments.update_one({"assignment_id": assignment_id},
                                        {"$set": {"status": "completed", "completed_at": now_iso()}})
    await log_event(assignment_id, user["user_id"], "quiz_submitted", None,
                    {"score": score, "passed": passed})
    return {"score": score, "passed": passed, "pass_score": pass_score, "answers": graded}


# ---------- REPORTS ----------
@api_router.get("/reports/overview")
async def reports_overview(admin=Depends(require_admin)):
    total_users = await db.users.count_documents({"role": "employee"})
    total_trainings = await db.trainings.count_documents({})
    total_assignments = await db.assignments.count_documents({})
    completed = await db.assignments.count_documents({"status": "completed"})
    in_progress = await db.assignments.count_documents({"status": {"$in": ["in_progress", "video_completed"]}})
    total_questions = await db.questions.count_documents({})
    scores = []
    async for p in db.progress.find({"quiz_attempts.0": {"$exists": True}}, {"_id": 0, "quiz_attempts": 1}):
        scores.append(p["quiz_attempts"][-1]["score"])
    trainings = await db.trainings.find({}, {"_id": 0}).to_list(500)
    per_training = []
    for t in trainings:
        cnt = await db.assignments.count_documents({"training_id": t["training_id"]})
        cmp = await db.assignments.count_documents({"training_id": t["training_id"], "status": "completed"})
        if cnt:
            per_training.append({"title": t["title"], "training_id": t["training_id"],
                                 "assigned": cnt, "completed": cmp})
    return {"total_users": total_users, "total_trainings": total_trainings,
            "total_assignments": total_assignments, "completed": completed,
            "in_progress": in_progress, "total_questions": total_questions,
            "completion_rate": round(completed / total_assignments * 100) if total_assignments else 0,
            "avg_quiz_score": round(sum(scores) / len(scores)) if scores else None,
            "per_training": per_training}


@api_router.get("/reports/trainings/{training_id}")
async def training_report(training_id: str, admin=Depends(require_admin)):
    t = await db.trainings.find_one({"training_id": training_id}, {"_id": 0})
    if not t:
        raise HTTPException(status_code=404, detail="Eğitim bulunamadı")
    assignments = await db.assignments.find({"training_id": training_id}, {"_id": 0}).to_list(1000)
    users = {u["user_id"]: u for u in await db.users.find({}, {"_id": 0}).to_list(1000)}
    rows = []
    duration = t.get("duration") or 0
    for a in assignments:
        prog = await db.progress.find_one({"assignment_id": a["assignment_id"]}, {"_id": 0})
        u = users.get(a["user_id"], {})
        last_quiz = (prog or {}).get("quiz_attempts", [])
        rows.append({
            "assignment_id": a["assignment_id"], "user_name": u.get("name", "-"),
            "user_email": u.get("email", "-"), "status": a["status"],
            "start_at": a["start_at"], "due_at": a.get("due_at"), "completed_at": a.get("completed_at"),
            "watch_pct": round(min(100, prog["max_position"] / duration * 100)) if prog and duration else 0,
            "watched_seconds": round(prog["watched_seconds"]) if prog else 0,
            "checkpoints_passed": len(prog["checkpoints_passed"]) if prog else 0,
            "checkpoints_total": len(t.get("checkpoints", [])),
            "checkpoint_fails": len([x for x in (prog or {}).get("checkpoint_attempts", []) if not x["passed"]]),
            "quiz_score": last_quiz[-1]["score"] if last_quiz else None,
            "quiz_attempts": len(last_quiz),
        })
    return {"training": {"title": t["title"], "duration": duration}, "rows": rows}


@api_router.get("/reports/trainings/{training_id}/export")
async def export_training_report(training_id: str, fmt: str = "pdf", admin=Depends(require_admin)):
    data = await training_report(training_id, admin)
    title = data["training"]["title"]
    rows = data["rows"]
    headers = ["Kullanıcı", "E-posta", "Durum", "İzleme %", "İzleme Süresi", "Kontrol N.", "Hata", "Sınav %", "Tamamlanma"]
    status_tr = {"assigned": "Atandı", "in_progress": "Devam Ediyor", "video_completed": "Video Bitti", "completed": "Tamamlandı"}

    def fmt_sec(s):
        return f"{int(s // 60):02d}:{int(s % 60):02d}"

    def fmt_d(iso):
        return iso[:10] if iso else "-"

    table_rows = [[r["user_name"], r["user_email"], status_tr.get(r["status"], r["status"]),
                   f"%{r['watch_pct']}", fmt_sec(r["watched_seconds"]),
                   f"{r['checkpoints_passed']}/{r['checkpoints_total']}", str(r["checkpoint_fails"]),
                   f"%{r['quiz_score']}" if r["quiz_score"] is not None else "-",
                   fmt_d(r.get("completed_at"))] for r in rows]

    slug = unicodedata.normalize("NFKD", title).encode("ascii", "ignore").decode().replace(" ", "_") or "rapor"
    date_str = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    gen_at = datetime.now(timezone.utc).strftime("%d.%m.%Y %H:%M UTC")

    if fmt == "excel":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook()
        ws = wb.active
        ws.title = "Eğitim Raporu"
        ws.append([f"Eğitim Raporu — {title}"])
        ws.append([f"Oluşturulma: {gen_at} · {len(rows)} atama"])
        ws.append([])
        ws.append(headers)
        for tr in table_rows:
            ws.append(tr)
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"].font = Font(size=9, color="888888")
        for cell in ws[4]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1A1A1A")
            cell.alignment = Alignment(horizontal="left")
        widths = [22, 30, 14, 10, 13, 10, 8, 9, 13]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = w
        ws.freeze_panes = "A5"
        buf = BytesIO()
        wb.save(buf)
        return Response(content=buf.getvalue(),
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f'attachment; filename="rapor_{slug}_{date_str}.xlsx"'})

    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import ParagraphStyle
    font_dir = Path("/usr/share/fonts/truetype/dejavu")
    if "DejaVu" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DejaVu", str(font_dir / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(font_dir / "DejaVuSans-Bold.ttf")))
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm, title=f"Eğitim Raporu — {title}")
    story = [
        Paragraph("Plena LMS — Eğitim Denetim Raporu", ParagraphStyle("h1", fontName="DejaVu-Bold", fontSize=16, leading=20)),
        Spacer(1, 2 * mm),
        Paragraph(f"Eğitim: {title}", ParagraphStyle("h2", fontName="DejaVu", fontSize=11, leading=14)),
        Paragraph(f"Oluşturulma: {gen_at} · {len(rows)} atama", ParagraphStyle("meta", fontName="DejaVu", fontSize=8, textColor=colors.HexColor("#888888"), leading=11)),
        Spacer(1, 6 * mm),
    ]
    cell_style = ParagraphStyle("cell", fontName="DejaVu", fontSize=8, leading=10)
    body = [[Paragraph(str(c), cell_style) for c in tr] for tr in table_rows]
    tbl = Table([headers] + body, repeatRows=1,
                colWidths=[42 * mm, 58 * mm, 26 * mm, 18 * mm, 24 * mm, 20 * mm, 14 * mm, 17 * mm, 24 * mm])
    tbl.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), "DejaVu-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A1A1A")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F7F5")]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#DDDDDD")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(tbl)
    doc.build(story)
    return Response(content=buf.getvalue(), media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="rapor_{slug}_{date_str}.pdf"'})


@api_router.get("/reports/assignments/{assignment_id}/detail")
async def assignment_detail(assignment_id: str, admin=Depends(require_admin)):
    a = await db.assignments.find_one({"assignment_id": assignment_id}, {"_id": 0})
    if not a:
        raise HTTPException(status_code=404, detail="Atama bulunamadı")
    prog = await db.progress.find_one({"assignment_id": assignment_id}, {"_id": 0})
    events = await db.events.find({"assignment_id": assignment_id}, {"_id": 0}).sort("created_at", 1).to_list(2000)
    user = await db.users.find_one({"user_id": a["user_id"]}, {"_id": 0})
    training = await db.trainings.find_one({"training_id": a["training_id"]}, {"_id": 0})
    return {"assignment": a, "progress": prog, "events": events,
            "user": user, "training_title": training["title"] if training else "-"}


@api_router.get("/emails")
async def list_emails(admin=Depends(require_admin)):
    return await db.email_logs.find({}, {"_id": 0}).sort("created_at", -1).to_list(500)


# ---------- REMINDER LOOP ----------
async def reminder_loop():
    while True:
        try:
            now = datetime.now(timezone.utc)
            cursor = db.assignments.find({"status": {"$ne": "completed"}, "reminder_days": {"$gt": 0},
                                          "start_at": {"$lte": now.isoformat()}}, {"_id": 0})
            async for a in cursor:
                base = parse_dt(a["last_reminder_at"] or a["start_at"])
                if now >= base + timedelta(days=a["reminder_days"]):
                    user = await db.users.find_one({"user_id": a["user_id"]}, {"_id": 0})
                    training = await db.trainings.find_one({"training_id": a["training_id"]}, {"_id": 0})
                    if user and training:
                        await log_email(user["email"], f"Hatırlatma: {training['title']}",
                                        f"Merhaba {user['name']}, '{training['title']}' eğitiminiz hâlâ tamamlanmadı. Lütfen en kısa sürede tamamlayın.",
                                        "reminder")
                        await db.assignments.update_one({"assignment_id": a["assignment_id"]},
                                                        {"$set": {"last_reminder_at": now.isoformat()}})
        except Exception as e:
            logger.error(f"Reminder loop error: {e}")
        await asyncio.sleep(1800)


@app.on_event("startup")
async def startup():
    existing = await db.users.find_one({"email": ADMIN_SEED_EMAIL})
    if not existing:
        await db.users.insert_one({"user_id": new_id("user"), "email": ADMIN_SEED_EMAIL,
                                   "name": "Tuğberk Kalay", "role": "admin", "status": "invited",
                                   "picture": None, "created_at": now_iso()})
    else:
        await db.users.update_one({"email": ADMIN_SEED_EMAIL}, {"$set": {"role": "admin"}})
    asyncio.create_task(reminder_loop())


app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
